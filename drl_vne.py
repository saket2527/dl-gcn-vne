"""Five-stage DRL policy for SAGIN virtual network embedding.

The policy follows the paper's extraction, convolution, probability, filtering,
and output stages. Because this repository's virtual-link delay demands (1-4)
are on a different scale from some SAGIN substrate delays (up to 34), delay is
used as a policy feature and path tie-breaker, not as a hard rejection limit.
The output ranking also discounts candidates that require more bandwidth-
weighted hops to already-mapped virtual neighbors, directly improving the
revenue-to-cost objective while preserving the learned policy's preferences.
"""

from __future__ import annotations

import argparse
import copy
import pickle
import random
from collections import deque
from datetime import timedelta
from pathlib import Path
from time import perf_counter
from typing import Any

import numpy as np
import torch
from openpyxl import load_workbook
from torch import Tensor, nn

DEFAULT_DRL_CHECKPOINT = Path(__file__).with_name("drl_vne.pth")
DEFAULT_SUBSTRATE = Path(__file__).with_name("1_sagin.pickle")
DEFAULT_WORKBOOK = Path(__file__).with_name("Results.xlsx")
DRL_ALGORITHM_NAME = "DRL-VNE"
ROUTE_COST_WEIGHT = 10.0


def _node_id(node: Any) -> int:
    return int(node)


def _edge_weight(graph: Any, left: int, right: int) -> float:
    for key in ((str(left), str(right)), (left, right)):
        if key in graph.edge_weights:
            return float(graph.edge_weights[key])
    raise KeyError(f"No bandwidth entry for substrate edge ({left}, {right})")


def _set_edge_weight(graph: Any, left: int, right: int, value: float) -> None:
    for key in ((str(left), str(right)), (left, right)):
        if key in graph.edge_weights:
            graph.edge_weights[key] = value
            return
    raise KeyError(f"No bandwidth entry for substrate edge ({left}, {right})")


def _edge_delay(graph: Any, left: int, right: int) -> float:
    for key in ((str(left), str(right)), (left, right)):
        if hasattr(graph, "delay") and key in graph.delay:
            return float(graph.delay[key])
    return 1.0


def undirected_edges(graph: Any):
    seen: set[tuple[int, int]] = set()
    for left, right in graph.edge_weights:
        edge = tuple(sorted((_node_id(left), _node_id(right))))
        if edge not in seen:
            seen.add(edge)
            yield edge


class DRLPolicy(nn.Module):
    """Five-stage policy: features, affine scoring, softmax, filter, ranking."""

    feature_count = 4

    def __init__(self) -> None:
        super().__init__()
        self.convolution = nn.Linear(self.feature_count, 1)

    def forward(self, features: Tensor, feasible_mask: Tensor | None = None) -> Tensor:
        scores = self.convolution(features).squeeze(-1)
        probabilities = torch.softmax(scores, dim=0)
        if feasible_mask is not None:
            probabilities = probabilities * feasible_mask.to(
                device=probabilities.device, dtype=probabilities.dtype
            )
            probabilities = probabilities / probabilities.sum().clamp_min(1e-12)
        return probabilities


class DRLVNEEnvironment:
    """SAGIN substrate with resource-aware features and transactional mapping."""

    def __init__(self, substrate: Any) -> None:
        self.substrate = substrate
        self.initial_cpu = {int(node): float(value) for node, value in substrate.node_weights.items()}
        self.initial_bandwidth = {
            (left, right): _edge_weight(substrate, left, right)
            for left, right in undirected_edges(substrate)
        }
        self._distances = self._all_pairs_hop_distances()

    def _all_pairs_hop_distances(self) -> list[list[float]]:
        count = self.substrate.nodes
        distances = [[float("inf")] * count for _ in range(count)]
        for source in range(count):
            distances[source][source] = 0.0
            queue = deque([source])
            while queue:
                node = queue.popleft()
                for neighbor in self.substrate.neighbours.get(node, set()):
                    neighbor = int(neighbor)
                    if distances[source][neighbor] == float("inf"):
                        distances[source][neighbor] = distances[source][node] + 1
                        queue.append(neighbor)
        return distances

    def feature_matrix(self, used_nodes: set[int]) -> Tensor:
        count = self.substrate.nodes
        cpu = [float(self.substrate.node_weights[node]) for node in range(count)]
        bandwidth_sums = [0.0] * count
        delay_sums = [0.0] * count
        for left, right in undirected_edges(self.substrate):
            bandwidth = _edge_weight(self.substrate, left, right)
            delay = _edge_delay(self.substrate, left, right)
            bandwidth_sums[left] += bandwidth
            bandwidth_sums[right] += bandwidth
            delay_sums[left] += delay
            delay_sums[right] += delay

        available_nodes = [node for node in range(count) if node not in used_nodes]
        average_distances = []
        for node in range(count):
            destinations = [other for other in available_nodes if other != node]
            reachable = [self._distances[node][other] for other in destinations]
            reachable = [distance for distance in reachable if distance != float("inf")]
            average_distances.append(
                sum(reachable) / len(reachable) if reachable else 0.0
            )

        columns = [cpu, bandwidth_sums, delay_sums, average_distances]
        normalized = []
        for values in columns:
            scale = max(values, default=0.0)
            normalized.append([value / scale if scale > 0 else 0.0 for value in values])
        return torch.tensor(list(zip(*normalized)), dtype=torch.float32)

    def shortest_feasible_path(self, source: int, target: int, demand: float) -> list[int]:
        if source == target:
            return [source]
        parent: dict[int, int | None] = {source: None}
        queue = deque([source])
        while queue:
            node = queue.popleft()
            neighbors = sorted(
                self.substrate.neighbours.get(node, set()),
                key=lambda neighbor: (
                    _edge_delay(self.substrate, node, int(neighbor)),
                    int(neighbor),
                ),
            )
            for neighbor_value in neighbors:
                neighbor = int(neighbor_value)
                if neighbor in parent or _edge_weight(self.substrate, node, neighbor) < demand:
                    continue
                parent[neighbor] = node
                if neighbor == target:
                    path = [target]
                    while parent[path[-1]] is not None:
                        path.append(parent[path[-1]])
                    return list(reversed(path))
                queue.append(neighbor)
        return []

    @staticmethod
    def virtual_edges(request: Any) -> list[tuple[int, int, float]]:
        edges: dict[tuple[int, int], float] = {}
        for left, right in request.edge_weights:
            edge = tuple(sorted((int(left), int(right))))
            edges.setdefault(edge, _edge_weight(request, edge[0], edge[1]))
        return [(left, right, demand) for (left, right), demand in sorted(edges.items())]

    def feasible_hop_distances(self, target: int, demand: float) -> list[float]:
        distances = [float("inf")] * self.substrate.nodes
        distances[target] = 0.0
        queue = deque([target])
        while queue:
            node = queue.popleft()
            for neighbor_value in self.substrate.neighbours.get(node, set()):
                neighbor = int(neighbor_value)
                if distances[neighbor] != float("inf"):
                    continue
                if _edge_weight(self.substrate, node, neighbor) < demand:
                    continue
                distances[neighbor] = distances[node] + 1
                queue.append(neighbor)
        return distances

    def rollback(
        self,
        cpu_changes: list[tuple[int, float]],
        bandwidth_changes: list[tuple[int, int, float]],
    ) -> None:
        for node, demand in cpu_changes:
            self.substrate.node_weights[node] += demand
        for left, right, demand in bandwidth_changes:
            _set_edge_weight(
                self.substrate, left, right,
                _edge_weight(self.substrate, left, right) + demand,
            )
            _set_edge_weight(
                self.substrate, right, left,
                _edge_weight(self.substrate, right, left) + demand,
            )

    def embed_request(
        self,
        request: Any,
        policy: DRLPolicy,
        *,
        sample: bool,
    ) -> tuple[bool, dict[str, Any], list[Tensor], list[Tensor]]:
        cpu_changes: list[tuple[int, float]] = []
        bandwidth_changes: list[tuple[int, int, float]] = []
        placement: dict[int, int] = {}
        used_nodes: set[int] = set()
        log_probabilities: list[Tensor] = []
        entropies: list[Tensor] = []
        edges = self.virtual_edges(request)

        incident_edges: dict[int, list[float]] = {node: [] for node in range(request.nodes)}
        for left, right, demand in edges:
            incident_edges[left].append(demand)
            incident_edges[right].append(demand)
        order = sorted(
            range(request.nodes),
            key=lambda node: (
                len(incident_edges[node]),
                sum(incident_edges[node]),
                float(request.node_weights[node]),
            ),
            reverse=True,
        )

        for virtual_node in order:
            demand = float(request.node_weights[virtual_node])
            features = self.feature_matrix(used_nodes)
            mapped_neighbors = [
                (placement[right], bandwidth)
                for left, right, bandwidth in edges
                if left == virtual_node and right in placement
            ] + [
                (placement[left], bandwidth)
                for left, right, bandwidth in edges
                if right == virtual_node and left in placement
            ]
            mapped_distances = [
                (bandwidth, self.feasible_hop_distances(target, bandwidth))
                for target, bandwidth in mapped_neighbors
            ]
            feasible = torch.tensor(
                [
                    node not in used_nodes
                    and float(self.substrate.node_weights[node]) >= demand
                    and all(distances[node] != float("inf") for _, distances in mapped_distances)
                    for node in range(self.substrate.nodes)
                ],
                dtype=torch.bool,
            )
            if not bool(feasible.any()):
                self.rollback(cpu_changes, bandwidth_changes)
                return False, {
                    "reason": "no_feasible_substrate_node",
                    "placement": placement,
                    "routes": [],
                }, log_probabilities, entropies

            probabilities = policy(features, feasible)
            route_costs = torch.tensor(
                [
                    sum(bandwidth * distances[node] for bandwidth, distances in mapped_distances)
                    if feasible[node] else 0.0
                    for node in range(self.substrate.nodes)
                ],
                dtype=probabilities.dtype,
            )
            total_incident_bandwidth = sum(
                bandwidth for left, right, bandwidth in edges
                if left == virtual_node or right == virtual_node
            )
            route_scale = max(1.0, total_incident_bandwidth)
            output_logits = torch.log(probabilities.clamp_min(1e-12))
            output_logits = output_logits - ROUTE_COST_WEIGHT * route_costs / route_scale
            output_logits = output_logits.masked_fill(~feasible, -torch.inf)
            distribution = torch.distributions.Categorical(logits=output_logits)
            action = distribution.sample() if sample else torch.argmax(output_logits)
            log_probabilities.append(distribution.log_prob(action))
            entropies.append(distribution.entropy())
            substrate_node = int(action.item())
            self.substrate.node_weights[substrate_node] -= demand
            cpu_changes.append((substrate_node, demand))
            placement[virtual_node] = substrate_node
            used_nodes.add(substrate_node)

        routes: list[dict[str, Any]] = []
        total_hops = 0
        total_delay = 0.0
        for left, right, demand in edges:
            path = self.shortest_feasible_path(placement[left], placement[right], demand)
            if not path:
                self.rollback(cpu_changes, bandwidth_changes)
                return False, {
                    "reason": "no_bandwidth_feasible_path",
                    "placement": placement,
                    "routes": routes,
                }, log_probabilities, entropies
            routes.append({
                "virtual_edge": (left, right),
                "substrate_path": path,
                "bandwidth_demand": demand,
            })
            total_hops += len(path) - 1
            for source, target in zip(path, path[1:]):
                total_delay += _edge_delay(self.substrate, source, target)
                _set_edge_weight(
                    self.substrate, source, target,
                    _edge_weight(self.substrate, source, target) - demand,
                )
                _set_edge_weight(
                    self.substrate, target, source,
                    _edge_weight(self.substrate, target, source) - demand,
                )
                bandwidth_changes.append((source, target, demand))

        node_cost = sum(demand for _, demand in cpu_changes)
        edge_cost = sum(
            route["bandwidth_demand"] * (len(route["substrate_path"]) - 1)
            for route in routes
        )
        revenue = node_cost + sum(demand for _, _, demand in edges)
        return True, {
            "placement": placement,
            "routes": routes,
            "node_cost": node_cost,
            "edge_cost": edge_cost,
            "total_cost": node_cost + edge_cost,
            "revenue": revenue,
            "hops": total_hops,
            "total_delay": total_delay,
        }, log_probabilities, entropies


def _load_substrate(path: str | Path = DEFAULT_SUBSTRATE) -> Any:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"SAGIN topology pickle not found: {path}")
    with path.open("rb") as stream:
        return pickle.load(stream)


def _request_reward(success: bool, info: dict[str, Any]) -> float:
    if not success:
        return -1.0
    cost = float(info["total_cost"])
    return float(info["revenue"]) / max(cost, 1e-8)


def _create_requests(count: int, seed: int) -> list[Any]:
    from vne_u import create_vne

    random.seed(seed)
    requests = create_vne(no_requests=count)
    resource_rng = random.Random(seed)
    for request in requests:
        parameters = request.parameters
        for node in range(request.nodes):
            request.node_weights[node] = resource_rng.randrange(
                int(parameters.lower_node), int(parameters.upper_node)
            )
        edges = sorted({
            tuple(sorted((int(left), int(right))))
            for left, right in request.edge_weights
        })
        for left, right in edges:
            bandwidth = resource_rng.randrange(
                int(parameters.lower_edge), int(parameters.upper_edge)
            )
            delay = resource_rng.randrange(
                int(parameters.lower_delay), int(parameters.upper_delay)
            )
            for source, target in ((left, right), (right, left)):
                request.edge_weights[(str(source), str(target))] = bandwidth
                request.delay[(str(source), str(target))] = delay
    return requests


def train_policy(
    substrate: Any | None = None,
    *,
    substrate_path: str | Path = DEFAULT_SUBSTRATE,
    checkpoint: str | Path = DEFAULT_DRL_CHECKPOINT,
    training_requests: int = 1000,
    batch_size: int = 100,
    learning_rate: float = 1e-2,
    seed: int = 7,
) -> DRLPolicy:
    if training_requests < 1 or batch_size < 1:
        raise ValueError("training_requests and batch_size must be positive.")
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    substrate_template = substrate if substrate is not None else _load_substrate(substrate_path)
    policy = DRLPolicy()
    optimizer = torch.optim.Adam(policy.parameters(), lr=learning_rate)

    completed = 0
    batch_number = 0
    while completed < training_requests:
        current_batch_size = min(batch_size, training_requests - completed)
        environment = DRLVNEEnvironment(copy.deepcopy(substrate_template))
        rewards: list[float] = []
        request_log_probabilities: list[list[Tensor]] = []
        request_entropies: list[list[Tensor]] = []
        policy.train()
        with torch.enable_grad():
            for request in _create_requests(current_batch_size, seed + completed):
                success, info, log_probs, entropies = environment.embed_request(
                    request, policy, sample=True
                )
                rewards.append(_request_reward(success, info))
                request_log_probabilities.append(log_probs)
                request_entropies.append(entropies)
        nonempty = [
            index for index, log_probs in enumerate(request_log_probabilities) if log_probs
        ]
        if nonempty:
            reward_tensor = torch.tensor(rewards, dtype=torch.float32)
            baseline = reward_tensor.mean()
            losses = []
            for index in nonempty:
                log_prob_sum = torch.stack(request_log_probabilities[index]).sum()
                entropy = torch.stack(request_entropies[index]).mean()
                advantage = reward_tensor[index] - baseline
                losses.append(-advantage.detach() * log_prob_sum - 0.01 * entropy)
            loss = torch.stack(losses).mean()
            optimizer.zero_grad(set_to_none=True)
            loss.backward()
            torch.nn.utils.clip_grad_norm_(policy.parameters(), max_norm=5.0)
            optimizer.step()
        completed += current_batch_size
        batch_number += 1
        accepted = sum(reward >= 0 for reward in rewards)
        mean_reward = sum(rewards) / max(1, len(rewards))
        print(
            f"training batch={batch_number} requests={completed}/{training_requests} "
            f"accepted={accepted}/{current_batch_size} mean_reward={mean_reward:.4f}"
        )

    checkpoint_path = Path(checkpoint)
    checkpoint_path.parent.mkdir(parents=True, exist_ok=True)
    torch.save({
        "model_state_dict": policy.state_dict(),
        "metadata": {
            "method": "five-stage-policy-gradient",
            "feature_count": DRLPolicy.feature_count,
            "training_requests": training_requests,
            "batch_size": batch_size,
            "seed": seed,
        },
    }, checkpoint_path)
    return policy.eval()


def load_policy(checkpoint: str | Path = DEFAULT_DRL_CHECKPOINT) -> DRLPolicy:
    checkpoint_path = Path(checkpoint)
    if not checkpoint_path.exists():
        raise FileNotFoundError(
            f"DRL-VNE checkpoint not found: {checkpoint_path}. "
            "Train it with `python3 drl_vne.py --training-requests 1000`."
        )
    saved = torch.load(checkpoint_path, map_location="cpu")
    policy = DRLPolicy()
    policy.load_state_dict(saved["model_state_dict"])
    return policy.eval()


def evaluate(
    substrate: Any,
    vne_list: list[Any],
    checkpoint: str | Path = DEFAULT_DRL_CHECKPOINT,
) -> dict[str, Any]:
    if not vne_list:
        raise ValueError("At least one virtual network request is required.")
    policy = load_policy(checkpoint)
    original = copy.deepcopy(substrate)
    environment = DRLVNEEnvironment(copy.deepcopy(substrate))
    accepted = 0
    revenue = 0.0
    total_cost = 0.0
    total_hops = 0
    total_nodes = sum(request.nodes for request in vne_list)
    total_links = sum(len(DRLVNEEnvironment.virtual_edges(request)) for request in vne_list)
    start_time = perf_counter()
    for request in vne_list:
        success, info, _, _ = environment.embed_request(request, policy, sample=False)
        if success:
            accepted += 1
            revenue += info["revenue"]
            total_cost += info["total_cost"]
            total_hops += info["hops"]
    elapsed_seconds = perf_counter() - start_time

    substrate_edges = list(undirected_edges(original))
    used_nodes = sum(
        environment.substrate.node_weights[node] != original.node_weights[node]
        for node in range(original.nodes)
    )
    used_links = sum(
        _edge_weight(environment.substrate, left, right) != _edge_weight(original, left, right)
        for left, right in substrate_edges
    )
    node_utilization = [
        (float(original.node_weights[node]) - float(environment.substrate.node_weights[node]))
        / float(original.node_weights[node])
        for node in range(original.nodes)
        if environment.substrate.node_weights[node] != original.node_weights[node]
        and float(original.node_weights[node]) > 0
    ]
    link_utilization = [
        (_edge_weight(original, left, right) - _edge_weight(environment.substrate, left, right))
        / _edge_weight(original, left, right)
        for left, right in substrate_edges
        if _edge_weight(environment.substrate, left, right) != _edge_weight(original, left, right)
        and _edge_weight(original, left, right) > 0
    ]
    pre_resource = sum(float(value) for value in original.node_weights.values()) + sum(
        _edge_weight(original, left, right) for left, right in substrate_edges
    )
    post_resource = sum(float(value) for value in environment.substrate.node_weights.values()) + sum(
        _edge_weight(environment.substrate, left, right) for left, right in substrate_edges
    )
    return {
        "revenue": revenue,
        "total_cost": total_cost,
        "accepted": accepted,
        "total_request": len(vne_list),
        "pre_resource": pre_resource,
        "post_resource": post_resource,
        "No_of_Links_used": used_links,
        "No_of_Nodes_used": used_nodes,
        "avg_path": total_hops / max(1, total_links),
        "avg_exec": timedelta(seconds=elapsed_seconds),
        "total_nodes": total_nodes,
        "total_links": total_links,
        "avg_bw": 100 * sum(link_utilization) / max(1, len(link_utilization)),
        "avg_crb": 100 * sum(node_utilization) / max(1, len(node_utilization)),
        "avg_link": 100 * used_links / max(1, len(substrate_edges)),
        "avg_node": 100 * used_nodes / max(1, original.nodes),
    }


def main(
    substrate: Any | None = None,
    vne_list: list[Any] | None = None,
    checkpoint: str | Path = DEFAULT_DRL_CHECKPOINT,
    device: str | torch.device | None = None,
) -> dict[str, Any]:
    """Evaluate the saved paper-style policy using the shared metrics contract."""
    del device
    if substrate is None or vne_list is None:
        import config
        substrate = substrate if substrate is not None else getattr(config, "substrate", None)
        vne_list = vne_list if vne_list is not None else getattr(config, "vne_list", None)
    if substrate is None:
        substrate = _load_substrate()
    if not vne_list:
        vne_list = _create_requests(5, 7)
    return evaluate(substrate, vne_list, checkpoint)


def write_result_to_workbook(
    result: dict[str, Any],
    workbook_path: str | Path = DEFAULT_WORKBOOK,
) -> int:
    workbook_file = Path(workbook_path)
    if not workbook_file.exists():
        raise FileNotFoundError(f"Results workbook not found: {workbook_file}")
    workbook = load_workbook(workbook_file)
    sheet = workbook.active
    headers = {
        str(sheet.cell(row=1, column=column).value): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(row=1, column=column).value is not None
    }
    required = {
        "algorithm", "revenue", "total_cost", "revenuetocostratio", "accepted",
        "total_request", "embeddingratio", "pre_resource", "post_resource", "consumed",
        "avg_bw", "avg_crb", "avg_link", "No_of_Links_used", "avg_node",
        "No_of_Nodes_used", "avg_path", "avg_exec", "total_nodes", "total_links",
    }
    missing = required - headers.keys()
    if missing:
        workbook.close()
        raise ValueError(f"Results workbook is missing columns: {', '.join(sorted(missing))}")

    algorithm_column = headers["algorithm"]
    paper_rows = [
        row for row in range(2, sheet.max_row + 1)
        if str(sheet.cell(row=row, column=algorithm_column).value).strip().upper() == DRL_ALGORITHM_NAME
    ]
    if paper_rows:
        row_index = paper_rows[-1]
    else:
        algorithm_rows = [
            row for row in range(2, sheet.max_row + 1)
            if sheet.cell(row=row, column=algorithm_column).value is not None
        ]
        row_index = max(algorithm_rows, default=1) + 1
        sheet.insert_rows(row_index)

    source_row = row_index - 1
    while source_row > 1 and not any(
        sheet.cell(row=source_row, column=column).value is not None
        for column in range(1, sheet.max_column + 1)
    ):
        source_row -= 1
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(row=source_row, column=column)
        target = sheet.cell(row=row_index, column=column)
        if source.has_style:
            from copy import copy as copy_style
            target._style = copy_style(source._style)
        if source.number_format:
            target.number_format = source.number_format
    sheet.row_dimensions[row_index].height = sheet.row_dimensions[source_row].height

    column_values = {
        "algorithm": DRL_ALGORITHM_NAME,
        "revenue": result["revenue"],
        "total_cost": result["total_cost"],
        "revenuetocostratio": 100 * result["revenue"] / result["total_cost"] if result["total_cost"] else 0.0,
        "accepted": result["accepted"],
        "total_request": result["total_request"],
        "embeddingratio": 100 * result["accepted"] / result["total_request"],
        "pre_resource": result["pre_resource"],
        "post_resource": result["post_resource"],
        "consumed": result["pre_resource"] - result["post_resource"],
        "avg_bw": result["avg_bw"],
        "avg_crb": result["avg_crb"],
        "avg_link": result["avg_link"],
        "No_of_Links_used": result["No_of_Links_used"],
        "avg_node": result["avg_node"],
        "No_of_Nodes_used": result["No_of_Nodes_used"],
        "avg_path": result["avg_path"],
        "avg_exec": result["avg_exec"].total_seconds() * 1000 / max(1, result["total_request"]),
        "total_nodes": result["total_nodes"],
        "total_links": result["total_links"],
    }
    for header, value in column_values.items():
        sheet.cell(row=row_index, column=headers[header], value=value)
    if sheet.cell(row=1, column=1).value is None:
        numeric_indices = [
            sheet.cell(row=row, column=1).value
            for row in range(2, sheet.max_row + 1)
            if isinstance(sheet.cell(row=row, column=1).value, int)
        ]
        sheet.cell(row=row_index, column=1, value=max(numeric_indices, default=-1) + 1)
    workbook.save(workbook_file)
    workbook.close()
    return row_index


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--substrate-pickle", type=Path, default=DEFAULT_SUBSTRATE)
    parser.add_argument("--checkpoint", type=Path, default=DEFAULT_DRL_CHECKPOINT)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--training-requests", type=int, default=1000)
    parser.add_argument("--batch-size", type=int, default=100)
    parser.add_argument("--evaluation-requests", type=int, default=5)
    parser.add_argument("--seed", type=int, default=7)
    parser.add_argument("--skip-training", action="store_true")
    return parser.parse_args()


if __name__ == "__main__":
    arguments = _parse_args()
    substrate = _load_substrate(arguments.substrate_pickle)
    if not arguments.skip_training:
        train_policy(
            substrate,
            checkpoint=arguments.checkpoint,
            training_requests=arguments.training_requests,
            batch_size=arguments.batch_size,
            seed=arguments.seed,
        )
    np.random.seed(arguments.seed)
    torch.manual_seed(arguments.seed)
    requests = _create_requests(arguments.evaluation_requests, arguments.seed + arguments.training_requests)
    metrics = evaluate(substrate, requests, arguments.checkpoint)
    row = write_result_to_workbook(metrics, arguments.workbook)
    print(
        f"Evaluation: {metrics['accepted']}/{metrics['total_request']} accepted, "
        f"revenue/cost={100 * metrics['revenue'] / metrics['total_cost'] if metrics['total_cost'] else 0:.2f}%. "
        f"Added {DRL_ALGORITHM_NAME} at Excel row {row} in {arguments.workbook}."
    )
